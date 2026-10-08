"""Bulk invoice upload: a spreadsheet of rows becomes draft proformas, all or nothing."""

import io
from datetime import datetime, timedelta

from openpyxl import Workbook, load_workbook

from app.models import AuditLog, Client, ClientInvoice, Service
from app.services.numbering import business_date, month_end
from app.services.agents import house_agent
from tests.conftest import DEFAULT_SERVICE_TITLE, auth_header

HEADER = "Client,Service,Amount,Description,Due date,TDS %,Interstate,Remarks\n"


def csv_file(*rows: str, header: str = HEADER) -> dict:
    body = header + "".join(r + "\n" for r in rows)
    return {"file": ("invoices.csv", body.encode("utf-8"), "text/csv")}


def xlsx_file(rows: list[list]) -> dict:
    wb = Workbook()
    ws = wb.active
    ws.append(["Client", "Service", "Amount", "Description", "Due date", "TDS %", "Interstate", "Remarks"])
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return {"file": ("invoices.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}


def preview(client, user, files):
    return client.post("/api/invoices/bulk/preview", files=files, headers=auth_header(user))


def commit(client, user, files):
    return client.post("/api/invoices/bulk", files=files, headers=auth_header(user))


def second_client(db, agent_id=None):
    row = Client(business_name="Bharat Steel", email="ap@bharat.example", gst_number="29CCCCC2222C1Z5",
                 agent_id=agent_id or house_agent(db).id)
    db.add(row)
    db.commit()
    db.refresh(row)
    return row


def test_preview_totals_rows_for_several_clients(client, users, client_row, db):
    second_client(db)
    res = preview(client, users["exec"], csv_file(
        f"Acme Traders,{DEFAULT_SERVICE_TITLE},10000,,,,,",
        f"29CCCCC2222C1Z5,{DEFAULT_SERVICE_TITLE},\"1,00,000\",Q3 retainer,,,yes,",
    ))
    assert res.status_code == 200, res.text
    p = res.json()
    assert p["ready"] is True
    assert p["invoice_count"] == 2
    assert p["client_count"] == 2
    assert p["taxable_total"] == 110000
    assert p["gst_total"] == 19800  # the service is 18%
    assert p["grand_total"] == 129800
    first, second = p["rows"]
    assert first["row"] == 2 and first["client_name"] == "Acme Traders" and first["errors"] == []
    assert second["client_name"] == "Bharat Steel"  # matched by GSTIN
    assert second["is_interstate"] is True
    assert second["description"] == "Q3 retainer"


def test_client_and_service_names_match_ignoring_case_and_spaces(client, users, client_row):
    res = preview(client, users["exec"], csv_file(f"  acme TRADERS ,{DEFAULT_SERVICE_TITLE.upper()},500,,,,,"))
    assert res.json()["ready"] is True


def test_every_problem_is_reported_against_its_row(client, users, client_row, db):
    db.add(Service(title="Old service", gst_rate=18, is_active=False))
    db.commit()
    yesterday = (business_date() - timedelta(days=1)).isoformat()
    res = preview(client, users["exec"], csv_file(
        f"Relience Industries,{DEFAULT_SERVICE_TITLE},1000,,,,,",
        "Acme Traders,Tax filing,1000,,,,,",
        "Acme Traders,Old service,1000,,,,,",
        f"Acme Traders,{DEFAULT_SERVICE_TITLE},abc,,,,,",
        f"Acme Traders,{DEFAULT_SERVICE_TITLE},0,,,,,",
        f"Acme Traders,{DEFAULT_SERVICE_TITLE},1000,,{yesterday},,,",
        f"Acme Traders,{DEFAULT_SERVICE_TITLE},1000,,31-02-2026,,,",
        f"Acme Traders,{DEFAULT_SERVICE_TITLE},1000,,,150,,",
        f"Acme Traders,{DEFAULT_SERVICE_TITLE},1000,,,,maybe,",
        f",{DEFAULT_SERVICE_TITLE},1000,,,,,",
    ))
    p = res.json()
    assert p["ready"] is False
    assert p["error_count"] == 10
    errors = {r["row"]: " ".join(r["errors"]) for r in p["rows"]}
    assert "No client called 'Relience Industries'" in errors[2]
    assert "No service called 'Tax filing'" in errors[3]
    assert "deactivated" in errors[4]
    assert "Amount" in errors[5]
    assert "more than zero" in errors[6]
    assert "in the past" in errors[7]
    assert "Due date" in errors[8]
    assert "TDS" in errors[9]
    assert "Interstate" in errors[10]
    assert "Client is required" in errors[11]


def test_a_name_shared_by_two_clients_needs_the_gstin(client, users, client_row, db):
    db.add(Client(business_name="Acme Traders", email="other@acme.example", gst_number="07DDDDD3333D1Z5",
                  agent_id=house_agent(db).id))
    db.commit()
    p = preview(client, users["exec"], csv_file(f"Acme Traders,{DEFAULT_SERVICE_TITLE},1000,,,,,")).json()
    assert "use their GSTIN" in p["rows"][0]["errors"][0]
    p = preview(client, users["exec"], csv_file(f"27BBBBB1111B1Z5,{DEFAULT_SERVICE_TITLE},1000,,,,,")).json()
    assert p["ready"] is True


def test_repeated_rows_are_flagged_but_allowed(client, users, client_row):
    row = f"Acme Traders,{DEFAULT_SERVICE_TITLE},1000,,,,,"
    p = preview(client, users["exec"], csv_file(row, row)).json()
    assert p["ready"] is True
    assert "Same client, service and amount as row 2" in p["rows"][1]["warnings"][0]


def test_upload_creates_draft_proformas_in_row_order(client, users, client_row, db):
    agent_client = second_client(db)
    res = commit(client, users["exec"], csv_file(
        f"Acme Traders,{DEFAULT_SERVICE_TITLE},10000,,,,,",
        f"Bharat Steel,{DEFAULT_SERVICE_TITLE},2000,Setup fee,,2,,Per email",
    ))
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["created"] == 2
    numbers = [i["proforma_number"] for i in body["invoices"]]
    assert numbers[0].endswith("/0001") and numbers[1].endswith("/0002")

    invoices = db.query(ClientInvoice).order_by(ClientInvoice.id).all()
    assert [i.status.value for i in invoices] == ["Draft", "Draft"]
    assert invoices[0].invoice_number is None  # a tax invoice number comes only at issue
    assert invoices[0].due_date == month_end(business_date())
    assert invoices[0].total_amount == 11800
    assert invoices[1].client_id == agent_client.id
    assert invoices[1].agent_id == agent_client.agent_id
    assert invoices[1].service_description == "Setup fee"
    assert invoices[1].tds_rate == 2
    assert invoices[1].internal_remarks == "Per email"
    assert invoices[1].created_by_id == users["exec"].id
    assert db.query(AuditLog).filter(AuditLog.action == "Bulk uploaded proformas").count() == 1


def test_one_bad_row_means_nothing_is_created(client, users, client_row, db):
    res = commit(client, users["exec"], csv_file(
        f"Acme Traders,{DEFAULT_SERVICE_TITLE},10000,,,,,",
        f"Nobody Ltd,{DEFAULT_SERVICE_TITLE},10000,,,,,",
    ))
    assert res.status_code == 422
    assert "1 row" in res.json()["detail"]
    assert db.query(ClientInvoice).count() == 0

    # The proforma series was not consumed by the failed upload.
    good = commit(client, users["exec"], csv_file(f"Acme Traders,{DEFAULT_SERVICE_TITLE},10000,,,,,"))
    assert good.json()["invoices"][0]["proforma_number"].endswith("/0001")


def test_excel_upload_with_real_date_cells(client, users, client_row):
    due = month_end(business_date()) + timedelta(days=15)
    res = commit(client, users["exec"], xlsx_file([
        ["Acme Traders", DEFAULT_SERVICE_TITLE, 2500.5, None, datetime(due.year, due.month, due.day), 10, "No", None],
        [None, None, None, None, None, None, None, None],  # blank rows are ignored
    ]))
    assert res.status_code == 201, res.text
    assert res.json()["created"] == 1


def test_indian_style_dates_are_read_day_first(client, users, client_row):
    due = business_date() + timedelta(days=40)
    p = preview(client, users["exec"], csv_file(f"Acme Traders,{DEFAULT_SERVICE_TITLE},1000,,{due:%d/%m/%Y},,,")).json()
    assert p["rows"][0]["due_date"] == due.isoformat()


def test_file_shape_problems_are_rejected_up_front(client, users, client_row):
    res = preview(client, users["exec"], csv_file("Acme Traders,1000", header="Client,Amount\n"))
    assert res.status_code == 400
    assert "Service" in res.json()["detail"]

    res = preview(client, users["exec"], {"file": ("invoices.pdf", b"%PDF", "application/pdf")})
    assert res.status_code == 400

    res = preview(client, users["exec"], csv_file())
    assert res.status_code == 400
    assert "no invoice rows" in res.json()["detail"]

    many = [f"Acme Traders,{DEFAULT_SERVICE_TITLE},1000,,,,," for _ in range(501)]
    res = preview(client, users["exec"], csv_file(*many))
    assert res.status_code == 400
    assert "500" in res.json()["detail"]


def test_template_lists_active_clients_and_services(client, users, client_row):
    res = client.get("/api/invoices/bulk/template", headers=auth_header(users["exec"]))
    assert res.status_code == 200
    wb = load_workbook(io.BytesIO(res.content))
    headers = [c.value for c in wb["Invoices"][1]]
    assert headers[:3] == ["Client", "Service", "Amount"]
    lists = wb["Lists"]
    assert "Acme Traders" in [c.value for c in lists["A"]]
    assert DEFAULT_SERVICE_TITLE in [c.value for c in lists["B"]]


def test_bulk_upload_needs_a_finance_login(client):
    assert client.post("/api/invoices/bulk/preview", files=csv_file()).status_code == 401
